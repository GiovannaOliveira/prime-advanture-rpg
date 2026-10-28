"""Offline development importer. Requires openpyxl; never needed by the browser.
Usage: python3 scripts/import-pokemon.py path/to/database.xlsx
Only public species metadata and selected embedded images are exported.
No collection, storage, owner, capture-date, or trainer fields are retained.
"""
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import json,re,sys,posixpath,hashlib
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent.parent
SELECTED=[1,4,7,10,12,16,23,25,27,35,37,39,41,43,54,58,60,63,66,69,74,77,79,81,92,94,95,104,124,125,129,131,133,134,143,144,146,147,149,151]
NS={'xdr':'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
source=Path(sys.argv[1]);workbook=load_workbook(source,read_only=True,data_only=True)
rows=[r for r in workbook['Pokédex'].iter_rows(min_row=3,values_only=True) if isinstance(r[1],(int,float)) and r[6]]
by_dex={}
for row in rows:
 dex=int(row[1])
 if dex not in by_dex or (row[9]=='No' and by_dex[dex][9]!='No'):by_dex[dex]=row

def flag(v):return v not in [None,'','-','No','no','#N/A',False]
def metadata(row):
 return {'nationalDex':int(row[1]),'name':str(row[6]),'generation':str(row[4]),'region':str(row[3]),'types':[str(t) for t in row[7:9] if t not in [None,'','-','No']],'form':str(row[9]),'hasEvolution':row[10]=='Yes','legendary':flag(row[34]),'mythical':flag(row[35])}

# Canonical national species catalog is reference data, not 1,025 simultaneous encounters.
catalog=[metadata(by_dex[dex]) for dex in sorted(by_dex)]
(ROOT/'data/pokemon-catalog.json').write_text(json.dumps({'sourceFile':source.name,'sourceSha256':hashlib.sha256(source.read_bytes()).hexdigest(),'sourceRows':len(rows),'speciesCount':len(catalog),'species':catalog},ensure_ascii=False,indent=2)+'\n')
selected=[]
with ZipFile(source) as z:
 rels={r.attrib['Id']:r.attrib['Target'] for r in ET.fromstring(z.read('xl/_rels/cellimages.xml.rels')) if r.attrib.get('TargetMode')!='External'}
 images={}
 for item in ET.fromstring(z.read('xl/cellimages.xml')):
  name=item.find('.//xdr:cNvPr',NS);blip=item.find('.//a:blip',NS)
  if name is not None and blip is not None:
   rid=blip.attrib.get('{'+NS['r']+'}embed');target=rels.get(rid)
   if target:images[name.attrib['name']]=posixpath.normpath(posixpath.join('xl',target))
 for dex in SELECTED:
  row=by_dex[dex];info=metadata(row);match=re.search(r'DISPIMG\("([^\"]+)"',str(row[2]))
  assert match and match[1] in images,('Missing embedded sprite',dex,info['name'])
  member=images[match[1]];assert member.startswith('xl/media/') and member in z.namelist()
  suffix=Path(member).suffix.lower();assert suffix=='.png',(member,suffix)
  identifier=f'pokemon-{dex:04d}';relative=f'assets/creatures/{identifier}.png';payload=z.read(member)
  (ROOT/relative).write_bytes(payload)
  info.update({'id':identifier,'asset':relative,'imageSha256':hashlib.sha256(payload).hexdigest()})
  selected.append(info)
workbook.close()
# These stats are game-specific balancing values, not stats claimed from the spreadsheet.
rarities={1:'Uncommon',4:'Uncommon',7:'Uncommon',12:'Uncommon',25:'Uncommon',35:'Uncommon',37:'Uncommon',54:'Uncommon',58:'Uncommon',63:'Uncommon',66:'Uncommon',74:'Common',77:'Uncommon',79:'Uncommon',81:'Uncommon',92:'Uncommon',94:'Rare',95:'Rare',124:'Rare',125:'Rare',131:'Rare',133:'Uncommon',134:'Rare',143:'Rare',147:'Uncommon',149:'Rare'}
profiles={'Grass':(39,10,6),'Fire':(35,12,4),'Water':(41,9,6),'Electric':(34,12,4),'Bug':(32,9,4),'Normal':(40,9,5),'Poison':(36,11,5),'Ground':(43,10,7),'Fairy':(37,10,6),'Fighting':(40,12,5),'Rock':(46,10,8),'Psychic':(34,12,4),'Ghost':(33,12,4),'Ice':(39,11,5),'Dragon':(44,12,7),'Dark':(35,12,5),'Steel':(44,10,8),'Flying':(34,11,4)}
moves={'Grass':('Leaf Sweep','leafEffect'),'Fire':('Warm Flame','fireEffect'),'Water':('Water Splash','waterEffect'),'Electric':('Spark Flash','magicEffect'),'Bug':('Pollen Drift','leafEffect'),'Normal':('Quick Dash','shadowEffect'),'Poison':('Poison Mist','magicEffect'),'Ground':('Dust Swirl','shadowEffect'),'Fairy':('Fairy Gleam','magicEffect'),'Fighting':('Brave Strike','shadowEffect'),'Rock':('Rock Roll','shadowEffect'),'Psychic':('Mind Wave','magicEffect'),'Ghost':('Shadow Flicker','shadowEffect'),'Ice':('Ice Shimmer','iceEffect'),'Dragon':('Dragon Pulse','magicEffect'),'Dark':('Night Dash','shadowEffect'),'Steel':('Metal Flash','magicEffect'),'Flying':('Wing Gust','leafEffect')}
for p in selected:
 p['rarity']='Legendary' if p['legendary'] or p['mythical'] else rarities.get(p['nationalDex'],'Common')
 p['hp'],p['attack'],p['defense']=profiles[p['types'][0]]
 if p['rarity']=='Legendary':p['hp']+=12;p['attack']+=2;p['defense']+=2
 elif p['rarity']=='Rare':p['hp']+=5;p['attack']+=1
 p['special'],p['effect']=moves[p['types'][0]]
 p['description']=(f"{p['name']} has the {p['types'][0]} and {p['types'][1]} types." if len(p['types'])==2 else f"{p['name']} is {'an' if p['types'][0][0] in 'AEIOU' else 'a'} {p['types'][0]}-type Pokémon.")+" Explore its habitat to meet it."
 if p['legendary'] or p['mythical']:p['requiresNote']={144:'frost',146:'ember',151:'ruins'}.get(p['nationalDex'],'ruins')
script="'use strict';\n// Imported public metadata and unmodified embedded PNGs from the supplied workbook.\n// HP, attack, defense, rarity tiers and moves below are Prime Adventure balancing choices.\n(() => {\n  const entries = "+json.dumps(selected,ensure_ascii=False,indent=2)+";\n  for (const row of entries) {\n    Prime.data.assets[row.id] = row.asset;\n    Prime.data.creatures[row.id] = {...row, type:row.types[0], assetFormat:'portrait', origin:'pokemon-workbook'};\n  }\n})();\n"
(ROOT/'src/data/pokemon.js').write_text(script)
print(json.dumps({'recordsRead':len(rows),'catalogSpecies':len(catalog),'playableImportedSpecies':len(selected),'legendary':[p['name'] for p in selected if p['rarity']=='Legendary'],'personalFieldsExported':False},indent=2))
